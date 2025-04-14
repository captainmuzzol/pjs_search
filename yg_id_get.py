# --*-- coding: utf-8 --*--
# 提取原告的身份证号程序(公司不存在身份证号，故不提取)
# 作者：许钦滔
# 修改时间：20240407
import os
import re
import sys

from docx import Document
from openpyxl import Workbook
from PyQt5.QtCore import Qt
from PyQt5.QtWidgets import (QApplication, QFileDialog, QMessageBox,
                             QProgressBar)
from win32com import client


def extract_id_numbers(file_paths):
    id_numbers = []
    for file_path in file_paths:
        if file_path.endswith(".txt"):
            with open(file_path, encoding='utf-8' ,errors='ignore') as f:
                content = f.read()
        else: # .docx文件
            doc = Document(file_path)
            content = "\n".join([para.text for para in doc.paragraphs])
        matches = re.findall(r'(?:原告人|原告：)(.*?)，.*?(\d{17}(\d|X|x))', content)
        # id_numbers.extend(matches)
        for match in matches:
            name = match[0].strip()
            id_number = match[1]
            id_numbers.append((file_path, name, id_number))
    return id_numbers

def write_to_excel(results):
    wb = Workbook()
    ws = wb.active
    ws.append(["文件名称", "原告姓名", "身份证号"])
    for result in results:
        ws.append(result)
    desktop = os.path.join(os.path.expanduser("~"), "Desktop")
    output_file = os.path.join(desktop, "身份证提取结果.xlsx")
    wb.save(output_file)

def convert_doc_to_txt(folder_path, progress_bar):
    input_folder_path = folder_path
    output_folder_path = folder_path

    # 获取指定文件夹下的所有doc和docx文件
    doc_files = []
    for file_name in os.listdir(input_folder_path):
        if file_name.endswith(".doc") or file_name.endswith(".docx"):
            doc_files.append(os.path.join(input_folder_path, file_name))
    total_files = len(doc_files)

    # 将doc和docx文件转化为utf-8的txt文件并保存到指定文件夹下
    word = client.Dispatch("Word.Application")
    failed_files = []  # 用于存储转换失败的文件
    for i, doc_file in enumerate(doc_files):
        try:
            # 打开doc文件并转化为txt文件
            doc = word.Documents.Open(doc_file)
            txt_file = os.path.join(output_folder_path, os.path.splitext(os.path.basename(doc_file))[0] + ".txt")
            doc.SaveAs(txt_file, FileFormat=4)
            doc.Close()

            # 将txt文件转化为utf-8编码
            with open(txt_file, "r", encoding='gbk', errors='ignore') as f:
                content = f.read()
            with open(txt_file, "w", encoding="utf-8") as f:
                f.write(content)

            # 更新进度条
            progress_bar.setValue(int((i+1)/total_files*100))

        except Exception as e:
            print("文件{}转换失败，错误信息：{}".format(doc_file, e))
            failed_files.append(doc_file)

    word.Quit()

    # 打印转换失败的文件
    print("以下文件转换失败：")
    for file in failed_files:
        print(file)

def main():
    app2 = QApplication(sys.argv)

    # 弹出文件夹选择框
    folder_path = QFileDialog.getExistingDirectory(None, '选择文件夹')
    if not folder_path:
        QMessageBox.information(None, "结果", "你没有选择文件夹！")
        app2.quit()

    try:
        # 显示开始运算提示框
        QMessageBox.information(None, "提示", "1.点击OK开始运算，根据文件数量会影响计算速度，完成后会有提示，请耐心等待。\n 2.请确保目标路径下不存在.txt文件，不然可能会被删除！")

        # 创建进度条
        progress_bar = QProgressBar()
        progress_bar.setAlignment(Qt.AlignCenter)
        progress_bar.setWindowTitle("转换进度")
        progress_bar.show()

        # 将所有.doc文件转换为.txt文件
        convert_doc_to_txt(folder_path, progress_bar)

        # 隐藏进度条
        progress_bar.hide()

        # 读取文件夹下所有txt和.docx文件
        files = [os.path.join(folder_path, f) for f in os.listdir(folder_path) if f.endswith(".txt") or f.endswith(".docx")]

        # 提取身份证号
        id_numbers = extract_id_numbers(files)

        # 将结果写入Excel表格
        write_to_excel(id_numbers)

        # 显示结果弹窗
        QMessageBox.information(None, "结果", "身份证提取完成，结果已保存到桌面的身份证提取结果.xlsx文件中。")

        app2.quit()
    except FileNotFoundError:
        QMessageBox.information(None, "结果", "你没有选择文件夹！")
        app2.quit()

if __name__ == "__main__":
    main()